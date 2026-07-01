#!/usr/bin/env python3
"""Core injection engine. Fill template from paper content.
Configure SRC (template path), OUT_DIR, and content strings before running."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os, re
from datetime import datetime

SRC = 'template.docx'
OUT_DIR = './output'

doc = Document(SRC)

def add_run(para, text, font_name='SimSun', font_size=Pt(12), bold=False):
    run = para.add_run(text)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = font_size
    run.font.bold = bold
    return run

def set_content_fmt(para, line_spacing=Pt(22)):
    pf = para.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = line_spacing
    pf.first_line_indent = Cm(0.74)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

def tick_checkbox(cell, label_text):
    for para in cell.paragraphs:
        for i, run in enumerate(para.runs):
            if run.text.strip() == '\u25a1' and i+1 < len(para.runs):
                if label_text in para.runs[i+1].text:
                    run.text = '\u2611'
                    return True
    return False

def add_page_break_before(para):
    pPr = para._element.get_or_add_pPr()
    pb = OxmlElement('w:pageBreakBefore')
    pPr.append(pb)

def split_bold_line(line):
    s = line.strip()
    if not s: return None, line
    standalone = [r'^[123]\\.[^:]{0,30}$', r'^Expected']
    for pat in standalone:
        if re.match(pat, s): return s, ''
    if re.match(r'^Innovation\\s+[A-C]', s):
        m = re.match(r'^Innovation\\s+[A-C][:;]', s)
        return s[:m.end()], s[m.end():]
    for pat, end in [(r'^\\([a-g]\\)', ')'), (r'^\\d+[.]', '.')]:
        m = re.match(pat, s)
        if m: return s[:m.end()], s[m.end():]
    return None, line

def inject_section(paras, label_kw, text, page_break=False):
    label_idx = None
    for i, p in enumerate(paras):
        if label_kw in p.text and len(p.text.strip()) > 10:
            label_idx = i
            break
    if label_idx is None: return
    if page_break: add_page_break_before(paras[label_idx])
    body = doc.element.body
    children = list(body)
    label_elem = paras[label_idx]._element
    label_body_idx = children.index(label_elem)
    section_keywords = ['Section', 'Header', 'Next']
    end_idx = len(paras) - 1
    for j in range(label_idx + 1, len(paras)):
        t = paras[j].text.strip()
        if t and any(kw in t for kw in section_keywords):
            end_idx = j - 1
            break
    elems_to_remove = []
    for j in range(label_idx + 1, end_idx + 1):
        if j < len(paras):
            elems_to_remove.append(paras[j]._element)
    for elem in elems_to_remove:
        try: body.remove(elem)
        except: pass
    # Build new content paragraphs
    def new_content_para():
        p = OxmlElement('w:p')
        pPr = OxmlElement('w:pPr')
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'both'); pPr.append(jc)
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:line'), '440'); sp.set(qn('w:lineRule'), 'auto')
        sp.set(qn('w:before'), '0'); sp.set(qn('w:after'), '0')
        pPr.append(sp)
        ind = OxmlElement('w:ind'); ind.set(qn('w:firstLine'), '420'); pPr.append(ind)
        p.append(pPr)
        return p
    def add_text(p, text, bold=False):
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24'); rPr.append(sz)
        rf = OxmlElement('w:rFonts')
        rf.set(qn('w:ascii'), 'SimSun'); rf.set(qn('w:eastAsia'), 'SimSun')
        rPr.append(rf)
        if bold:
            b = OxmlElement('w:b'); rPr.append(b)
        r.append(rPr)
        t_elem = OxmlElement('w:t')
        t_elem.set(qn('xml:space'), 'preserve')
        t_elem.text = text
        r.append(t_elem)
        p.append(r)
    blocks = re.split(r'\\n\\n', text)
    new_elems = []
    for block in blocks:
        if not block.strip(): continue
        p = new_content_para()
        lines = block.split('\\n')
        for li, line in enumerate(lines):
            if not line.strip(): continue
            if li > 0:
                br = OxmlElement('w:r')
                br_e = OxmlElement('w:br')
                br.append(br_e)
                p.append(br)
            bold_part, normal_part = split_bold_line(line)
            if bold_part:
                add_text(p, bold_part, bold=True)
                if normal_part:
                    add_text(p, normal_part, bold=False)
            else:
                add_text(p, normal_part or line, bold=False)
        new_elems.append(p)
    insert_pos = label_body_idx + 1
    for elem in new_elems:
        body.insert(insert_pos, elem)
        insert_pos += 1
    return label_idx

# === Customize content below ===
intro_text = 'Your project introduction here. Use \\n\\n for paragraph breaks.'
tech_text = 'Your technical innovation content here.'
compare_text = 'Your comparison and evaluation here.'
app_text = 'Your application outlook here.'

# === Run injection ===
paras = doc.paragraphs
inject_section(paras, 'intro label keyword', intro_text)
inject_section(paras, 'tech label keyword', tech_text, page_break=True)
inject_section(paras, 'compare label keyword', compare_text, page_break=True)
inject_section(paras, 'app label keyword', app_text, page_break=True)

# === Save ===
ts = datetime.now().strftime('%Y%m%d')
existing = [f for f in os.listdir(OUT_DIR) if ts in f]
v = len(existing) + 1
dst = os.path.join(OUT_DIR, f'filled_{ts}_v{v}.docx')
doc.save(dst)
print(f'Saved: {dst}')
