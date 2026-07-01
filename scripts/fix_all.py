#!/usr/bin/env python3
"""One-shot post-processing fixer.
Usage: python fix_all.py <filled.docx> [--keep-pagebreaks] [--margins L R]

Fixes: empty paras, formatting, outline, margins, bold audit."""
import sys, os, glob, re
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def get_next_version(out_dir):
    existing = []
    for f in os.listdir(out_dir):
        m = re.search(r'_v(\d+)\.docx$', f)
        if m: existing.append(int(m.group(1)))
    return max(existing) + 1 if existing else 1

def fix_all(docx_path, keep_pagebreaks=False, custom_margins=None):
    doc = Document(docx_path)
    body = doc.element.body
    out_dir = os.path.dirname(docx_path)
    base = os.path.basename(docx_path).rsplit('_v', 1)[0]
    report = []
    # 1. Remove empty paragraphs
    removed = 0
    for p in list(doc.paragraphs):
        if not p.text.strip():
            has_content = any('drawing' in c.tag for c in (r._element for r in p.runs) for c in [c])
            if not has_content:
                try: body.remove(p._element); removed += 1
                except: pass
    report.append(f'Empty paras removed: {removed}')
    # 2. Standardize body text
    section_kw = ['基本情况','完成人','项目简介','详细内容','科技奖励','知识产权','发表论文','计划资助','附件目录','主要科技','第三方评价','推广应用','经济效益','书面材料']
    body_fixed = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t: continue
        is_label = any(kw in t for kw in section_kw) or re.match(r'^[一二三四五六七八九十]、', t)
        if is_label:
            p.paragraph_format.line_spacing = Pt(28)
        else:
            pf = p.paragraph_format
            if pf.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY: pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.line_spacing = Pt(22)
            pf.space_before = Pt(0); pf.space_after = Pt(0); pf.first_line_indent = Cm(0)
            for run in p.runs:
                if run.font.size and run.font.size != Pt(12): run.font.size = Pt(12)
                if run.font.name and run.font.name != 'SimSun':
                    run.font.name = 'SimSun'
                    if run._element.rPr is not None:
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
            body_fixed += 1
    report.append(f'Body standardized: {body_fixed}')
    # 3. Outline levels
    h = [0,0,0]
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t: continue
        pPr = p._element.get_or_add_pPr()
        for c in list(pPr):
            if c.tag.endswith('outlineLvl'): pPr.remove(c)
        if re.match(r'^[一二三四五六七八九]、', t):
            ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'),'0'); pPr.append(ol); h[0]+=1
        elif re.match(r'^[123]、', t):
            ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'),'1'); pPr.append(ol); h[1]+=1
        elif re.match(r'^创新点', t):
            ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'),'2'); pPr.append(ol); h[2]+=1
    report.append(f'Outline H1={h[0]} H2={h[1]} H3={h[2]}')
    # 4. Strip leading spaces
    stripped = sum(1 for p in doc.paragraphs if p.text and p.text != p.text.lstrip())
    for p in doc.paragraphs:
        if p.text: p.text = p.text.lstrip()
    report.append(f'Leading spaces stripped: {stripped}')
    # 5. Margins
    l, r = custom_margins if custom_margins else (Cm(2.9), Cm(2.9))
    for s in doc.sections: s.left_margin = l; s.right_margin = r
    # 6. Page breaks
    if not keep_pagebreaks:
        pb = 0
        for p in doc.paragraphs:
            pPr = p._element.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
            if pPr is not None:
                for c in list(pPr):
                    if c.tag.endswith('pageBreakBefore'): pPr.remove(c); pb += 1
        report.append(f'Page breaks removed: {pb}')
    # 7. Bold audit
    over = sum(1 for p in doc.paragraphs for r in p.runs if r.font.bold and (len(r.text)>100 or len([s for s in r.text.split('。') if s.strip()])>1))
    report.append(f'Bold audit: {"CLEAN" if over==0 else str(over)+" oversize"}')
    # Save
    v = get_next_version(out_dir)
    dst = os.path.join(out_dir, f'{base}_v{v}.docx')
    doc.save(dst)
    report.append(f'Saved: {os.path.basename(dst)}')
    return dst, report

if __name__ == '__main__':
    import argparse
    p = argparse.ArgumentParser()
    p.add_argument('docx')
    p.add_argument('--keep-pagebreaks',action='store_true')
    p.add_argument('--margins',nargs=2,type=float)
    a = p.parse_args()
    margins = tuple(Cm(m) for m in a.margins) if a.margins else None
    dst, report = fix_all(a.docx, a.keep_pagebreaks, margins)
    for r in report: print(f'  {r}')
