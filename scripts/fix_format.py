#!/usr/bin/env python3
"""Post-processing: remove empty paragraphs, standardize formatting."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os, glob, re
from datetime import datetime

def fix_format(docx_path):
    doc = Document(docx_path)
    body = doc.element.body
    out_dir = os.path.dirname(docx_path)
    # Remove empty paragraphs
    removed = 0
    for p in list(doc.paragraphs):
        if not p.text.strip():
            has_img = any('drawing' in c.tag for r in p.runs for c in r._element)
            if not has_img:
                try: body.remove(p._element); removed += 1
                except: pass
    print(f'Removed {removed} empty paragraphs')
    # Standardize body
    section_kw = ['基本情况','完成人','项目简介','详细内容','科技奖励','知识产权','发表论文','计划资助','附件','主要科技','第三方评价','推广应用','经济效益','书面材料']
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t: continue
        is_label = any(kw in t for kw in section_kw) or re.match(r'^[一二三四五六七八九十]、', t)
        if not is_label:
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.line_spacing = Pt(22)
            pf.space_before = Pt(0); pf.space_after = Pt(0)
            pf.first_line_indent = Cm(0)
            for run in p.runs:
                run.font.size = Pt(12)
                run.font.name = 'SimSun'
                if run._element.rPr is not None:
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
    # Add outline levels
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t: continue
        pPr = p._element.get_or_add_pPr()
        for c in list(pPr):
            if c.tag.endswith('outlineLvl'): pPr.remove(c)
        if re.match(r'^[一二三四五六七八九]、', t):
            ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'),'0'); pPr.append(ol)
        elif re.match(r'^[123]、', t):
            ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'),'1'); pPr.append(ol)
        elif re.match(r'^创新点', t):
            ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'),'2'); pPr.append(ol)
    # Save versioned
    ts = datetime.now().strftime('%Y%m%d')
    existing = [f for f in os.listdir(out_dir) if ts in f]
    v = len(existing) + 1
    base = os.path.basename(docx_path).rsplit('_v', 1)[0] if '_v' in docx_path else os.path.splitext(os.path.basename(docx_path))[0]
    dst = os.path.join(out_dir, f'{base}_v{v}.docx')
    doc.save(dst)
    print(f'Saved: {os.path.basename(dst)}')

if __name__ == '__main__':
    import sys
    if len(sys.argv) < 2:
        out_dir = '.'
        files = glob.glob(os.path.join(out_dir, '*_v*.docx'))
        if files:
            files.sort(key=os.path.getmtime, reverse=True)
            fix_format(files[0])
    else:
        fix_format(sys.argv[1])
