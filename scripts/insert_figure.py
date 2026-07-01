#!/usr/bin/env python3
"""Insert figure into latest versioned docx.
Usage: python insert_figure.py <image.png> <output_dir>"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os, glob, sys

if len(sys.argv) < 3:
    print('Usage: python insert_figure.py <image.png> <output_dir>')
    sys.exit(1)

IMG = sys.argv[1]
OUT_DIR = sys.argv[2]

files = [f for f in glob.glob(os.path.join(OUT_DIR, '*_v*.docx')) if not os.path.basename(f).startswith('~$')]
files.sort(key=os.path.getmtime, reverse=True)
if not files:
    print('No versioned docx found')
    sys.exit(1)

doc = Document(files[0])
print(f'Target: {os.path.basename(files[0])}')

# Find tech paragraph
paras = doc.paragraphs
tech_idx = None
for i, p in enumerate(paras):
    if '技术原理' in p.text or '1．' in p.text:
        tech_idx = i
        break

if tech_idx is None:
    print('ERROR: Cannot find tech content paragraph')
    sys.exit(1)

# Add image paragraph at end, then reposition
img_para = doc.add_paragraph()
img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
img_run = img_para.add_run()
img_run.add_picture(IMG, width=Inches(5.5))

# Caption
cap_para = doc.add_paragraph()
cap_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
cap_para.paragraph_format.line_spacing = Pt(22)
cap_para.paragraph_format.first_line_indent = Cm(0.74)
cap_run = cap_para.add_run('Figure caption - edit me')
cap_run.font.size = Pt(12)
cap_run.font.name = 'SimSun'

# Reposition after tech paragraph
body = doc.element.body
children = list(body)
tech_elem = paras[tech_idx]._element
tech_body_idx = children.index(tech_elem)
img_elem = img_para._element
cap_elem = cap_para._element
body.remove(img_elem)
body.remove(cap_elem)
body.insert(tech_body_idx + 1, img_elem)
body.insert(tech_body_idx + 2, cap_elem)

doc.save(files[0])
print('Done')
